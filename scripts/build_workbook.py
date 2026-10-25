"""
build_workbook.py
-----------------
Builds workbook/Precast_Cost_Analysis.xlsx from the CSV files in data/raw/.

Every calculated cell is a live Excel formula, so the workbook recalculates
when an input changes. Lookups are written with XLOOKUP (Excel 2021 or
Microsoft 365). Use --compat to write INDEX/MATCH instead, for older Excel
versions or LibreOffice.

Usage:
    python scripts/build_workbook.py            # XLOOKUP version
    python scripts/build_workbook.py --compat   # INDEX/MATCH version
"""

from __future__ import annotations

import argparse
import csv
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "workbook" / "Precast_Cost_Analysis.xlsx"
MAX_ROW = 500  # formula ranges cover rows 2..500 so new rows are included

# ---------------------------------------------------------------- styles
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")      # hardcoded inputs
F_LINK = Font(name=FONT, size=10, color="008000")       # links to other sheets
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F3A5F")
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_RAW = PatternFill("solid", fgColor="1F3A5F")       # raw data header
FILL_CALC = PatternFill("solid", fgColor="5B7F95")      # calculated header
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_FLAG = PatternFill("solid", fgColor="F8CBAD")
FILL_OK = PatternFill("solid", fgColor="E2EFDA")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FMT_CUR2 = '$#,##0.00;($#,##0.00);"-"'
FMT_CUR0 = '$#,##0;($#,##0);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_NUM = '#,##0.####'
FMT_INT = '#,##0'


def read_csv(name: str) -> list[dict]:
    with (RAW / name).open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def num(value: str):
    """Convert a CSV string to int/float; keep blanks as None."""
    if value is None or value.strip() == "":
        return None
    v = float(value)
    return int(v) if v.is_integer() else v


def header(ws, row: int, cols: list[tuple[str, int, bool]]) -> None:
    """cols = (title, width, is_calculated)"""
    for i, (title, width, calc) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.font = F_HEAD
        c.fill = FILL_CALC if calc else FILL_RAW
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    longest = max(len(t) for t, _, _ in cols)
    ws.row_dimensions[row].height = 45 if longest > 18 else 30


def put(ws, row, col, value, font=F_BASE, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    c.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def title(ws, text: str, sub: str) -> None:
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB


def flag_status(ws, col_letter: str, last_row: int, first_row: int = 2) -> None:
    rng = f"{col_letter}{first_row}:{col_letter}{last_row}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'], fill=FILL_FLAG))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK))


def add_name(wb, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def rng(sheet: str, col: str) -> str:
    """Absolute bounded range on another sheet, e.g. Components!$I$2:$I$500"""
    return f"{sheet}!${col}$2:${col}${MAX_ROW}"


# ======================================================================
# INDEX(return_range, MATCH(key, lookup_range, 0))  ->  XLOOKUP(key, lookup_range, return_range)
_INDEX_MATCH = re.compile(r"INDEX\(([^,()]+),MATCH\(([^,()]+),([^,()]+),0\)\)")
# IFERROR(XLOOKUP(k, l, r), "text")  ->  XLOOKUP(k, l, r, "text")  (if_not_found argument)
_IFERROR_XL = re.compile(r'IFERROR\(_xlfn\.XLOOKUP\(([^,()]+),([^,()]+),([^,()]+)\),("[^"]*")\)')


def to_xlookup(formula: str) -> str:
    """Rewrite INDEX/MATCH lookups as XLOOKUP. Other formulas are returned unchanged.

    The _xlfn. prefix is how Excel stores functions added after Excel 2007; Excel hides it.
    """
    out = _INDEX_MATCH.sub(r"_xlfn.XLOOKUP(\2,\3,\1)", formula)
    return _IFERROR_XL.sub(r"_xlfn.XLOOKUP(\1,\2,\3,\4)", out)


def build(compat: bool = False) -> None:
    items = read_csv("items.csv")
    suppliers = read_csv("suppliers.csv")
    comps = read_csv("components.csv")
    bom = read_csv("bom.csv")
    sales = read_csv("sales_2025.csv")

    wb = Workbook()
    ws_about = wb.active
    ws_about.title = "About"
    ws_as = wb.create_sheet("Assumptions")
    ws_it = wb.create_sheet("Items")
    ws_su = wb.create_sheet("Suppliers")
    ws_co = wb.create_sheet("Components")
    ws_bo = wb.create_sheet("BOM")
    ws_sa = wb.create_sheet("Sales")
    ws_cr = wb.create_sheet("Cost_Rollup")
    ws_so = wb.create_sheet("Sourcing")
    ws_sm = wb.create_sheet("Summary")
    ws_dq = wb.create_sheet("Data_Quality")
    ws_oq = wb.create_sheet("Open_Questions")

    # ------------------------------------------------------------ Assumptions
    ws = ws_as
    title(ws, "Assumptions and Settings",
          "Yellow cells are inputs. Change them here and every sheet updates.")
    ws.column_dimensions["A"].width = 4
    for col, w in zip("BCDEF", (28, 14, 14, 12, 52)):
        ws.column_dimensions[col].width = w

    ws["B4"] = "Exchange rates to CAD"; ws["B4"].font = F_BOLD
    for i, h in enumerate(("Currency", "Rate to CAD", "Note"), start=2):
        c = ws.cell(row=5, column=i, value=h); c.font = F_HEAD; c.fill = FILL_CALC
    fx = [("CAD", 1.00, "Base currency"),
          ("USD", 1.37, "Assumed average 2025 rate (illustrative)"),
          ("EUR", 1.50, "Assumed average 2025 rate (illustrative)")]
    for r, (cur, rate, note) in enumerate(fx, start=6):
        put(ws, r, 2, cur, F_INPUT)
        c = put(ws, r, 3, rate, F_INPUT, "0.0000"); c.fill = FILL_ASSUME
        put(ws, r, 4, note, F_SUB)
    add_name(wb, "FX_Currency", "Assumptions!$B$6:$B$8")
    add_name(wb, "FX_Rate", "Assumptions!$C$6:$C$8")

    ws["B11"] = "Unit conversions"; ws["B11"].font = F_BOLD
    for i, h in enumerate(("Key (formula)", "From unit", "To unit", "Factor", "Note"), start=2):
        c = ws.cell(row=12, column=i, value=h); c.font = F_HEAD; c.fill = FILL_CALC
    units = [("LB", "KG", 0.45359237, "1 lb = 0.45359237 kg"),
             ("G", "KG", 0.001, "1 g = 0.001 kg"),
             ("KG", "T", 0.001, "1 kg = 0.001 tonne")]
    for r, (fu, tu, f, note) in enumerate(units, start=13):
        put(ws, r, 2, f'=C{r}&">"&D{r}')
        put(ws, r, 3, fu, F_INPUT)
        put(ws, r, 4, tu, F_INPUT)
        c = put(ws, r, 5, f, F_INPUT, "0.########"); c.fill = FILL_ASSUME
        put(ws, r, 6, note, F_SUB)
    add_name(wb, "Unit_Key", "Assumptions!$B$13:$B$15")
    add_name(wb, "Unit_Factor", "Assumptions!$E$13:$E$15")

    ws["B18"] = "Thresholds and settings"; ws["B18"].font = F_BOLD
    settings = [("Home country code", "CA", None, "Used to split material cost into Domestic / Imported", "Home_Country"),
                ("Sales discount flag (vs list price)", 0.20, FMT_PCT, "Sales lines priced lower than this are flagged", "Discount_Threshold"),
                ("Long lead time flag (days)", 30, FMT_INT, "Suppliers above this lead time are flagged", "Lead_Time_Flag")]
    for r, (label, val, fmt, note, name) in enumerate(settings, start=19):
        put(ws, r, 2, label)
        c = put(ws, r, 3, val, F_INPUT, fmt); c.fill = FILL_ASSUME
        put(ws, r, 4, note, F_SUB)
        add_name(wb, name, f"Assumptions!$C${r}")

    ws["B24"] = "Key assumptions"; ws["B24"].font = F_BOLD
    notes = [
        "All data in this workbook is synthetic (invented) and is used for learning only.",
        "Unit costs are latest purchase prices, converted to CAD with the rates above.",
        "Conversion cost (labour + plant overhead) per item is an estimate held on the Items sheet.",
        "Concrete quantities use one simplified mix design per cubic metre for all products.",
        "If a component code appears twice, lookups use the FIRST row. Duplicates are flagged on Components.",
        "BOM lines with a missing cost are left out of material cost and flagged as 'Incomplete'.",
        "HS codes are illustrative and must be confirmed before any real use.",
        "Average selling price uses only 2025 invoice lines that have a price.",
    ]
    for r, text in enumerate(notes, start=25):
        ws.cell(row=r, column=2, value=f"A{r-24}.").font = F_BOLD
        ws.cell(row=r, column=3, value=text).font = F_BASE

    # ------------------------------------------------------------ Items
    ws = ws_it
    header(ws, 1, [("item_no", 14, False), ("description", 42, False), ("category", 13, False),
                   ("uom", 6, False), ("hs_code", 10, False), ("list_price_cad", 13, False),
                   ("conversion_cost_cad", 14, False), ("item_key", 14, True),
                   ("issues", 30, True), ("status", 9, True)])
    for r, it in enumerate(items, start=2):
        put(ws, r, 1, it["item_no"], F_INPUT)
        put(ws, r, 2, it["description"], F_INPUT)
        put(ws, r, 3, it["category"], F_INPUT)
        put(ws, r, 4, it["uom"], F_INPUT)
        put(ws, r, 5, it["hs_code"] or None, F_INPUT)
        put(ws, r, 6, num(it["list_price_cad"]), F_INPUT, FMT_CUR2)
        put(ws, r, 7, num(it["conversion_cost_cad"]), F_INPUT, FMT_CUR2)
        put(ws, r, 8, f"=UPPER(TRIM(A{r}))")
        put(ws, r, 9, f'=IF(OR(E{r}="",LEN(E{r})<>7,MID(E{r},5,1)<>"."),"HS code missing/format; ","")'
                      f'&IF(COUNTIF($H$2:$H${MAX_ROW},H{r})>1,"Duplicate item; ","")')
        put(ws, r, 10, f'=IF(I{r}="","OK","CHECK")')
    it_last = len(items) + 1
    flag_status(ws, "J", it_last)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:J{it_last}"

    # ------------------------------------------------------------ Suppliers
    ws = ws_su
    header(ws, 1, [("supplier_id", 12, False), ("supplier_name", 28, False), ("country", 9, False),
                   ("currency", 9, False), ("lead_time_days", 13, False), ("lead_time_flag", 16, True)])
    for r, s in enumerate(suppliers, start=2):
        put(ws, r, 1, s["supplier_id"], F_INPUT)
        put(ws, r, 2, s["supplier_name"], F_INPUT)
        put(ws, r, 3, s["country"], F_INPUT)
        put(ws, r, 4, s["currency"], F_INPUT)
        put(ws, r, 5, num(s["lead_time_days"]), F_INPUT, FMT_INT)
        put(ws, r, 6, f'=IF(E{r}>Lead_Time_Flag,"Long lead time","")')
    su_last = len(suppliers) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:F{su_last}"

    # ------------------------------------------------------------ Components
    ws = ws_co
    header(ws, 1, [("comp_code", 11, False), ("description", 34, False), ("uom", 6, False),
                   ("unit_cost", 10, False), ("currency", 9, False), ("supplier_id", 11, False),
                   ("origin_country", 10, False), ("hs_code", 10, False), ("comp_key", 11, True),
                   ("unit_cost_cad", 12, True), ("supplier_name", 26, True), ("key_count", 9, True),
                   ("issues", 40, True), ("status", 9, True)])
    for r, cp in enumerate(comps, start=2):
        put(ws, r, 1, cp["comp_code"], F_INPUT)
        put(ws, r, 2, cp["description"], F_INPUT)
        put(ws, r, 3, cp["uom"], F_INPUT)
        put(ws, r, 4, num(cp["unit_cost"]), F_INPUT, FMT_CUR2)
        put(ws, r, 5, cp["currency"], F_INPUT)
        put(ws, r, 6, cp["supplier_id"], F_INPUT)
        put(ws, r, 7, cp["origin_country"] or None, F_INPUT)
        put(ws, r, 8, cp["hs_code"] or None, F_INPUT)
        put(ws, r, 9, f"=UPPER(TRIM(A{r}))")
        put(ws, r, 10, f'=IF(D{r}="","",IFERROR(D{r}*INDEX(FX_Rate,MATCH(E{r},FX_Currency,0)),"NO FX RATE"))', fmt=FMT_CUR2)
        put(ws, r, 11, f'=IFERROR(INDEX({rng("Suppliers","B")},MATCH(F{r},{rng("Suppliers","A")},0)),"NOT FOUND")')
        put(ws, r, 12, f"=COUNTIF($I$2:$I${MAX_ROW},I{r})")
        put(ws, r, 13, f'=IF(D{r}="","Missing cost; ","")&IF(G{r}="","Missing origin; ","")'
                       f'&IF(K{r}="NOT FOUND","Unknown supplier; ","")&IF(L{r}>1,"Duplicate code; ","")'
                       f'&IF(OR(H{r}="",LEN(H{r})<>7,MID(H{r},5,1)<>"."),"HS code format; ","")'
                       f'&IF(J{r}="NO FX RATE","No FX rate; ","")')
        put(ws, r, 14, f'=IF(M{r}="","OK","CHECK")')
    co_last = len(comps) + 1
    flag_status(ws, "N", co_last)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:N{co_last}"

    # ------------------------------------------------------------ BOM
    ws = ws_bo
    header(ws, 1, [("parent_item", 13, False), ("line_no", 8, False), ("comp_code", 11, False),
                   ("qty_per", 9, False), ("uom", 6, False), ("comp_key", 11, True),
                   ("comp_description", 30, True), ("comp_uom", 9, True), ("conv_factor", 10, True),
                   ("qty_in_comp_uom", 12, True), ("unit_cost_cad", 12, True), ("extended_cost_cad", 13, True),
                   ("origin_country", 10, True), ("origin_group", 11, True), ("supplier_id", 11, True),
                   ("units_sold_2025", 11, True), ("spend_2025_cad", 13, True),
                   ("issues", 34, True), ("status", 9, True)])
    m = f'MATCH(F{{r}},{rng("Components","I")},0)'  # template, filled per row
    for r, b in enumerate(bom, start=2):
        mr = m.format(r=r)
        put(ws, r, 1, b["parent_item"], F_INPUT)
        put(ws, r, 2, num(b["line_no"]), F_INPUT)
        put(ws, r, 3, b["comp_code"], F_INPUT)
        put(ws, r, 4, num(b["qty_per"]), F_INPUT, FMT_NUM)
        put(ws, r, 5, b["uom"], F_INPUT)
        put(ws, r, 6, f"=UPPER(TRIM(C{r}))")
        put(ws, r, 7, f'=IFERROR(INDEX({rng("Components","B")},{mr}),"NOT FOUND")')
        put(ws, r, 8, f'=IF(G{r}="NOT FOUND","",INDEX({rng("Components","C")},{mr}))')
        put(ws, r, 9, f'=IF(G{r}="NOT FOUND","",IF(E{r}=H{r},1,IFERROR(INDEX(Unit_Factor,MATCH(E{r}&">"&H{r},Unit_Key,0)),"NO FACTOR")))', fmt="0.####")
        put(ws, r, 10, f'=IF(ISNUMBER(I{r}),D{r}*I{r},"")', fmt=FMT_NUM)
        put(ws, r, 11, f'=IF(G{r}="NOT FOUND","",INDEX({rng("Components","J")},{mr}))', fmt=FMT_CUR2)
        put(ws, r, 12, f'=IF(AND(ISNUMBER(J{r}),ISNUMBER(K{r})),J{r}*K{r},"")', fmt=FMT_CUR2)
        put(ws, r, 13, f'=IF(G{r}="NOT FOUND","",IF(INDEX({rng("Components","G")},{mr})="","UNKNOWN",INDEX({rng("Components","G")},{mr})))')
        put(ws, r, 14, f'=IF(M{r}="","",IF(M{r}="UNKNOWN","Unknown",IF(M{r}=Home_Country,"Domestic","Imported")))')
        put(ws, r, 15, f'=IF(G{r}="NOT FOUND","",INDEX({rng("Components","F")},{mr}))')
        put(ws, r, 16, f'=SUMIFS({rng("Sales","G")},{rng("Sales","I")},A{r})', fmt=FMT_INT)
        put(ws, r, 17, f'=IF(ISNUMBER(L{r}),L{r}*P{r},"")', fmt=FMT_CUR0)
        put(ws, r, 18, f'=IF(G{r}="NOT FOUND","Unknown component; ","")&IF(I{r}="NO FACTOR","No unit conversion; ","")'
                       f'&IF(AND(G{r}<>"NOT FOUND",K{r}=""),"Missing cost; ","")&IF(M{r}="UNKNOWN","Missing origin; ","")')
        put(ws, r, 19, f'=IF(R{r}="","OK","CHECK")')
    bo_last = len(bom) + 1
    flag_status(ws, "S", bo_last)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:S{bo_last}"

    # ------------------------------------------------------------ Sales
    ws = ws_sa
    header(ws, 1, [("invoice_no", 11, False), ("invoice_date", 11, False), ("customer_id", 11, False),
                   ("ship_to_region", 9, False), ("ship_to_country", 9, False), ("item_no", 13, False),
                   ("qty", 7, False), ("unit_price_cad", 12, False), ("item_key", 13, True),
                   ("category", 12, True), ("list_price_cad", 11, True), ("revenue_cad", 12, True),
                   ("priced_qty", 9, True), ("discount_vs_list", 10, True), ("unit_cost_cad", 11, True),
                   ("cost_of_sales_cad", 13, True), ("issues", 30, True), ("status", 9, True)])
    cr_first, cr_last = 5, 4 + len(items)
    for r, s in enumerate(sales, start=2):
        put(ws, r, 1, s["invoice_no"], F_INPUT)
        put(ws, r, 2, date.fromisoformat(s["invoice_date"]), F_INPUT, "yyyy-mm-dd")
        put(ws, r, 3, s["customer_id"], F_INPUT)
        put(ws, r, 4, s["ship_to_region"], F_INPUT)
        put(ws, r, 5, s["ship_to_country"], F_INPUT)
        put(ws, r, 6, s["item_no"], F_INPUT)
        put(ws, r, 7, num(s["qty"]), F_INPUT, FMT_INT)
        put(ws, r, 8, num(s["unit_price_cad"]), F_INPUT, FMT_CUR2)
        put(ws, r, 9, f"=UPPER(TRIM(F{r}))")
        put(ws, r, 10, f'=IFERROR(INDEX({rng("Items","C")},MATCH(I{r},{rng("Items","H")},0)),"UNKNOWN")')
        put(ws, r, 11, f'=IF(J{r}="UNKNOWN","",INDEX({rng("Items","F")},MATCH(I{r},{rng("Items","H")},0)))', fmt=FMT_CUR2)
        put(ws, r, 12, f'=IF(ISNUMBER(H{r}),G{r}*H{r},"")', fmt=FMT_CUR2)
        put(ws, r, 13, f'=IF(ISNUMBER(H{r}),G{r},0)', fmt=FMT_INT)
        put(ws, r, 14, f'=IF(AND(ISNUMBER(H{r}),ISNUMBER(K{r})),1-H{r}/K{r},"")', fmt=FMT_PCT)
        put(ws, r, 15, f'=IF(J{r}="UNKNOWN","",INDEX(Cost_Rollup!$F${cr_first}:$F${cr_last},'
                       f'MATCH(I{r},Cost_Rollup!$A${cr_first}:$A${cr_last},0)))', fmt=FMT_CUR2)
        put(ws, r, 16, f'=IF(AND(ISNUMBER(L{r}),ISNUMBER(O{r})),G{r}*O{r},"")', fmt=FMT_CUR2)
        put(ws, r, 17, f'=IF(H{r}="","Missing price; ","")&IF(J{r}="UNKNOWN","Unknown item; ","")'
                       f'&IF(AND(ISNUMBER(N{r}),N{r}>Discount_Threshold),"Discount above limit; ","")')
        put(ws, r, 18, f'=IF(Q{r}="","OK","CHECK")')
    sa_last = len(sales) + 1
    flag_status(ws, "R", sa_last)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:R{sa_last}"

    # ------------------------------------------------------------ Cost_Rollup
    ws = ws_cr
    title(ws, "Product Cost Roll-up",
          "Material cost from the BOM + estimated conversion cost, compared with 2025 average selling price.")
    header(ws, 4, [("Item", 13, True), ("Description", 38, True), ("Category", 12, True),
                   ("Material cost (CAD)", 12, True), ("Conversion cost (CAD)", 12, True),
                   ("Total unit cost (CAD)", 12, True), ("List price (CAD)", 11, True),
                   ("Avg selling price 2025 (CAD)", 13, True), ("Margin per unit (CAD)", 12, True),
                   ("Margin %", 9, True), ("Cost-to-price ratio", 10, True),
                   ("Domestic share of material cost", 12, True), ("BOM lines", 8, True),
                   ("Lines to review", 9, True), ("Data status", 12, True)])
    for i, r in enumerate(range(cr_first, cr_last + 1)):
        src = i + 2  # row on Items
        put(ws, r, 1, f"=Items!A{src}", F_LINK)
        put(ws, r, 2, f"=Items!B{src}", F_LINK)
        put(ws, r, 3, f"=Items!C{src}", F_LINK)
        put(ws, r, 4, f'=SUMIFS({rng("BOM","L")},{rng("BOM","A")},A{r})', fmt=FMT_CUR2)
        put(ws, r, 5, f"=Items!G{src}", F_LINK, FMT_CUR2)
        put(ws, r, 6, f"=D{r}+E{r}", fmt=FMT_CUR2)
        put(ws, r, 7, f"=Items!F{src}", F_LINK, FMT_CUR2)
        put(ws, r, 8, f'=IFERROR(SUMIFS({rng("Sales","L")},{rng("Sales","I")},A{r})'
                      f'/SUMIFS({rng("Sales","M")},{rng("Sales","I")},A{r}),"")', fmt=FMT_CUR2)
        put(ws, r, 9, f'=IF(ISNUMBER(H{r}),H{r}-F{r},"")', fmt=FMT_CUR2)
        put(ws, r, 10, f'=IF(ISNUMBER(H{r}),I{r}/H{r},"")', fmt=FMT_PCT)
        put(ws, r, 11, f'=IF(ISNUMBER(H{r}),F{r}/H{r},"")', fmt="0.00")
        put(ws, r, 12, f'=IF(D{r}>0,SUMIFS({rng("BOM","L")},{rng("BOM","A")},A{r},{rng("BOM","N")},"Domestic")/D{r},"")', fmt=FMT_PCT)
        put(ws, r, 13, f'=COUNTIF({rng("BOM","A")},A{r})', fmt=FMT_INT)
        put(ws, r, 14, f'=COUNTIFS({rng("BOM","A")},A{r},{rng("BOM","S")},"CHECK")', fmt=FMT_INT)
        put(ws, r, 15, f'=IF(N{r}>0,"Incomplete","Complete")')
    ws.conditional_formatting.add(f"O{cr_first}:O{cr_last}",
                                  CellIsRule(operator="equal", formula=['"Incomplete"'], fill=FILL_FLAG))
    ws.conditional_formatting.add(f"O{cr_first}:O{cr_last}",
                                  CellIsRule(operator="equal", formula=['"Complete"'], fill=FILL_OK))
    n = cr_last + 2
    ws.cell(row=n, column=1, value="Note: 'Incomplete' items have BOM lines to review (missing cost, origin or part number). "
                                    "A missing cost means material cost is understated. See BOM (status = CHECK).").font = F_SUB
    ws.freeze_panes = f"B{cr_first}"

    # ------------------------------------------------------------ Sourcing
    ws = ws_so
    title(ws, "Sourcing Analysis (2025)",
          "Material spend = BOM extended cost x units sold in 2025. Shows where purchased material comes from.")
    header(ws, 4, [("Origin country", 26, True), ("Material spend 2025 (CAD)", 16, True),
                   ("Share of spend", 11, True), ("BOM lines", 10, True), ("", 4, True),
                   ("Supplier", 11, True), ("Supplier name", 26, True), ("Country", 9, True),
                   ("Lead time (days)", 10, True), ("Material spend 2025 (CAD)", 16, True),
                   ("Share of spend", 11, True), ("Lead time flag", 15, True)])
    ws.cell(row=4, column=5).fill = PatternFill(fill_type=None)
    countries = ["CA", "US", "MX", "TR", "KR", "DE", "IN", "UNKNOWN"]
    c_first = 5
    c_last = c_first + len(countries) - 1
    tot_r = c_last + 1
    for i, code in enumerate(countries):
        r = c_first + i
        put(ws, r, 1, code, F_INPUT)
        put(ws, r, 2, f'=SUMIFS({rng("BOM","Q")},{rng("BOM","M")},A{r})', fmt=FMT_CUR0)
        put(ws, r, 3, f"=IF($B${tot_r}>0,B{r}/$B${tot_r},0)", fmt=FMT_PCT)
        put(ws, r, 4, f'=COUNTIF({rng("BOM","M")},A{r})', fmt=FMT_INT)
    put(ws, tot_r, 1, "Total costed spend", F_BOLD).fill = FILL_TOTAL
    put(ws, tot_r, 2, f"=SUM(B{c_first}:B{c_last})", F_BOLD, FMT_CUR0).fill = FILL_TOTAL
    put(ws, tot_r, 3, f"=SUM(C{c_first}:C{c_last})", F_BOLD, FMT_PCT).fill = FILL_TOTAL
    put(ws, tot_r, 4, f"=SUM(D{c_first}:D{c_last})", F_BOLD, FMT_INT).fill = FILL_TOTAL
    put(ws, tot_r + 1, 1, "Check: all BOM spend")
    put(ws, tot_r + 1, 2, f'=SUM({rng("BOM","Q")})', fmt=FMT_CUR0)
    put(ws, tot_r + 2, 1, "Difference (should be 0)")
    put(ws, tot_r + 2, 2, f"=B{tot_r+1}-B{tot_r}", fmt=FMT_CUR0)
    ws.conditional_formatting.add(f"B{tot_r+2}", FormulaRule(formula=[f"ABS(B{tot_r+2})>0.5"], fill=FILL_FLAG))
    ws.cell(row=tot_r + 3, column=1,
            value="If the difference is not 0, add the missing country code to column A.").font = F_SUB

    s_first = 5
    s_last = s_first + len(suppliers) - 1
    for i in range(len(suppliers)):
        r = s_first + i
        src = i + 2
        put(ws, r, 6, f"=Suppliers!A{src}", F_LINK)
        put(ws, r, 7, f"=Suppliers!B{src}", F_LINK)
        put(ws, r, 8, f"=Suppliers!C{src}", F_LINK)
        put(ws, r, 9, f"=Suppliers!E{src}", F_LINK, FMT_INT)
        put(ws, r, 10, f'=SUMIFS({rng("BOM","Q")},{rng("BOM","O")},F{r})', fmt=FMT_CUR0)
        put(ws, r, 11, f"=IF($B${tot_r}>0,J{r}/$B${tot_r},0)", fmt=FMT_PCT)
        put(ws, r, 12, f"=Suppliers!F{src}", F_LINK)
    ur = s_last + 1
    put(ws, ur, 6, "(other)")
    put(ws, ur, 7, "Not in supplier list")
    put(ws, ur, 10, f"=B{tot_r}-SUM(J{s_first}:J{s_last})", fmt=FMT_CUR0)
    put(ws, ur, 11, f"=IF($B${tot_r}>0,J{ur}/$B${tot_r},0)", fmt=FMT_PCT)
    ws.conditional_formatting.add(f"J{ur}", FormulaRule(formula=[f"J{ur}>0.5"], fill=FILL_FLAG))
    put(ws, ur + 1, 6, "Total", F_BOLD).fill = FILL_TOTAL
    put(ws, ur + 1, 10, f"=SUM(J{s_first}:J{ur})", F_BOLD, FMT_CUR0).fill = FILL_TOTAL
    put(ws, ur + 1, 11, f"=SUM(K{s_first}:K{ur})", F_BOLD, FMT_PCT).fill = FILL_TOTAL
    ws.conditional_formatting.add(f"L{s_first}:L{s_last}",
                                  CellIsRule(operator="equal", formula=['"Long lead time"'], fill=FILL_FLAG))

    # ------------------------------------------------------------ Summary
    ws = ws_sm
    title(ws, "Sales and Margin Summary (2025)",
          "Built with SUMIFS. A PivotTable version is one of the exercises in docs/EXERCISES.md.")
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18
    kpis = [
        ("Revenue, all priced lines (CAD)", f'=SUM({rng("Sales","L")})', FMT_CUR0),
        ("Revenue with known cost (CAD)", f'=SUMIFS({rng("Sales","L")},{rng("Sales","J")},"<>UNKNOWN")', FMT_CUR0),
        ("Cost of sales (CAD)", f'=SUM({rng("Sales","P")})', FMT_CUR0),
        ("Gross margin (CAD)", "=B5-B6", FMT_CUR0),
        ("Gross margin %", "=IF(B5>0,B7/B5,0)", FMT_PCT),
        ("Units shipped", f'=SUM({rng("Sales","G")})', FMT_INT),
        ("Invoice lines", f'=COUNTA({rng("Sales","A")})', FMT_INT),
        ("Invoice lines to review", f'=COUNTIF({rng("Sales","R")},"CHECK")', FMT_INT),
    ]
    ws["A3"] = "Key figures"; ws["A3"].font = F_BOLD
    for i, (label, f, fmt) in enumerate(kpis):
        r = 4 + i
        put(ws, r, 1, label)
        put(ws, r, 2, f, F_BOLD, fmt)

    r0 = 14
    ws.cell(row=r0 - 1, column=1, value="By product category").font = F_BOLD
    header_cols = ["Category", "Units", "Revenue (CAD)", "Cost of sales (CAD)", "Gross margin (CAD)", "GM %"]
    for i, h in enumerate(header_cols, start=1):
        c = ws.cell(row=r0, column=i, value=h); c.font = F_HEAD; c.fill = FILL_CALC; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    cats = ["Pipe", "Box Culvert", "Manhole", "Catch Basin", "Barrier", "Slab"]
    for i, cat in enumerate(cats):
        r = r0 + 1 + i
        put(ws, r, 1, cat, F_INPUT)
        put(ws, r, 2, f'=SUMIFS({rng("Sales","G")},{rng("Sales","J")},A{r})', fmt=FMT_INT)
        put(ws, r, 3, f'=SUMIFS({rng("Sales","L")},{rng("Sales","J")},A{r})', fmt=FMT_CUR0)
        put(ws, r, 4, f'=SUMIFS({rng("Sales","P")},{rng("Sales","J")},A{r})', fmt=FMT_CUR0)
        put(ws, r, 5, f"=C{r}-D{r}", fmt=FMT_CUR0)
        put(ws, r, 6, f"=IF(C{r}>0,E{r}/C{r},0)", fmt=FMT_PCT)
    ct = r0 + 1 + len(cats)
    put(ws, ct, 1, "Total", F_BOLD).fill = FILL_TOTAL
    for col, fmt in zip("BCDE", (FMT_INT, FMT_CUR0, FMT_CUR0, FMT_CUR0)):
        put(ws, ct, " BCDE".index(col) + 1, f"=SUM({col}{r0+1}:{col}{ct-1})", F_BOLD, fmt).fill = FILL_TOTAL
    put(ws, ct, 6, f"=IF(C{ct}>0,E{ct}/C{ct},0)", F_BOLD, FMT_PCT).fill = FILL_TOTAL
    ws.cell(row=ct + 1, column=1,
            value="Totals exclude invoice lines for unknown items. Items with incomplete BOM costs overstate margin.").font = F_SUB

    r1 = ct + 4
    ws.cell(row=r1 - 1, column=1, value="By ship-to country").font = F_BOLD
    for i, h in enumerate(["Country", "Units", "Revenue (CAD)", "Share of revenue"], start=1):
        c = ws.cell(row=r1, column=i, value=h); c.font = F_HEAD; c.fill = FILL_CALC; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, cc in enumerate(["CA", "US"]):
        r = r1 + 1 + i
        put(ws, r, 1, cc, F_INPUT)
        put(ws, r, 2, f'=SUMIFS({rng("Sales","G")},{rng("Sales","E")},A{r})', fmt=FMT_INT)
        put(ws, r, 3, f'=SUMIFS({rng("Sales","L")},{rng("Sales","E")},A{r})', fmt=FMT_CUR0)
        put(ws, r, 4, f"=IF($B$4>0,C{r}/$B$4,0)", fmt=FMT_PCT)

    # ------------------------------------------------------------ Data_Quality
    ws = ws_dq
    title(ws, "Data Quality Checks",
          "Each check counts flagged rows on its source sheet. Filter that sheet's 'status' column to CHECK to see them.")
    header(ws, 4, [("ID", 7, True), ("Sheet", 12, True), ("Check", 40, True), ("How it is tested", 52, True),
                   ("Rows flagged", 10, True), ("Severity", 10, True), ("Suggested action", 46, True)])
    bo_rng = lambda c: f"BOM!${c}$2:${c}${MAX_ROW}"  # noqa: E731
    checks = [
        ("Items", "HS code missing or wrong format", "Blank, not 7 characters, or no dot in position 5 (####.##)",
         f'=COUNTIF({rng("Items","I")},"*HS code*")', "Medium", "Confirm HS code with the product team"),
        ("Components", "Unit cost missing", "unit_cost is blank",
         f'=COUNTIF({rng("Components","M")},"*Missing cost*")', "High", "Get latest purchase price from purchasing"),
        ("Components", "Country of origin missing", "origin_country is blank",
         f'=COUNTIF({rng("Components","M")},"*Missing origin*")', "High", "Request origin information from the supplier"),
        ("Components", "Supplier not in supplier list", "supplier_id has no match on Suppliers",
         f'=COUNTIF({rng("Components","M")},"*Unknown supplier*")', "Medium", "Add the supplier or correct the ID"),
        ("Components", "Duplicate component code", "Same code after TRIM + UPPER appears more than once",
         f'=COUNTIF({rng("Components","M")},"*Duplicate*")', "High", "Keep one record; confirm which cost is current"),
        ("Components", "HS code format", "Blank, not 7 characters, or no dot in position 5",
         f'=COUNTIF({rng("Components","M")},"*HS code*")', "Low", "Correct the format (####.##)"),
        ("BOM", "Component code not found", "comp_key has no match on Components",
         f'=COUNTIF({bo_rng("R")},"*Unknown component*")', "High", "Ask engineering for the correct part number"),
        ("BOM", "No unit conversion available", "BOM unit differs from component unit and no factor exists",
         f'=COUNTIF({bo_rng("R")},"*No unit conversion*")', "High", "Add the factor on Assumptions"),
        ("BOM", "Line has no cost", "Component found but its unit cost is blank",
         f'=COUNTIF({bo_rng("R")},"*Missing cost*")', "High", "Fix the component cost (see Components)"),
        ("BOM", "Line has no country of origin", "Component found but origin is blank",
         f'=COUNTIF({bo_rng("R")},"*Missing origin*")', "Medium", "Fix the component origin (see Components)"),
        ("BOM", "Unit converted automatically (info)", "BOM unit differs from component unit; factor applied",
         f'=SUMPRODUCT(({bo_rng("E")}<>{bo_rng("H")})*({bo_rng("H")}<>""))', "Info",
         "Confirm the BOM unit is intended (e.g. LB vs KG)"),
        ("BOM", "Code cleaned by TRIM/UPPER (info)", "Original code differs from cleaned key (case or spaces)",
         f'=SUMPRODUCT(({bo_rng("C")}<>"")*NOT(EXACT({bo_rng("C")},{bo_rng("F")})))', "Info",
         "Standardise the code in the source system"),
        ("Sales", "Unit price missing", "unit_price_cad is blank",
         f'=COUNTIF({rng("Sales","Q")},"*Missing price*")', "High", "Check the invoice in the ERP"),
        ("Sales", "Item not in item list", "item_no has no match on Items",
         f'=COUNTIF({rng("Sales","Q")},"*Unknown item*")', "High", "Add the item or correct the item number"),
        ("Sales", "Discount above limit", "Price is more than the Assumptions threshold below list price",
         f'=COUNTIF({rng("Sales","Q")},"*Discount*")', "Medium", "Confirm the discount was approved"),
    ]
    for i, (sheet, chk, how, f, sev, act) in enumerate(checks):
        r = 5 + i
        put(ws, r, 1, f"DQ-{i+1:02d}", F_BOLD)
        put(ws, r, 2, sheet)
        put(ws, r, 3, chk)
        put(ws, r, 4, how)
        put(ws, r, 5, f, fmt=FMT_INT)
        put(ws, r, 6, sev)
        put(ws, r, 7, act)
    dq_last = 4 + len(checks)
    ws.conditional_formatting.add(
        f"E5:E{dq_last}",
        FormulaRule(formula=[f'AND(E5>0,F5<>"Info")'], fill=FILL_FLAG))
    ws.conditional_formatting.add(f"E5:E{dq_last}", CellIsRule(operator="equal", formula=["0"], fill=FILL_OK))
    tr = dq_last + 2
    put(ws, tr, 3, "Rows with status CHECK, by sheet", F_BOLD)
    for i, (name, col) in enumerate([("Items", "J"), ("Components", "N"), ("BOM", "S"), ("Sales", "R")]):
        r = tr + 1 + i
        put(ws, r, 3, name)
        put(ws, r, 5, f'=COUNTIF({name}!${col}$2:${col}${MAX_ROW},"CHECK")', fmt=FMT_INT)
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ Open_Questions
    ws = ws_oq
    title(ws, "Open Questions Log",
          "Questions raised by the data checks. Update status and answer as information arrives.")
    header(ws, 4, [("Q#", 6, False), ("Date raised", 11, False), ("Related check", 10, False),
                   ("Item / component", 15, False), ("Question", 60, False), ("Ask", 14, False),
                   ("Status", 9, False), ("Answer / notes", 40, False)])
    raised = date(2026, 10, 21)
    questions = [
        ("DQ-03", "GSK-1200", "What is the country of origin for the 1200 mm gasket? Request a supplier declaration.", "Purchasing"),
        ("DQ-02", "SEAL-BUT", "What is the current purchase price for butyl joint sealant (per metre)?", "Purchasing"),
        ("DQ-04", "CST-FRM", "Supplier SUP-11 is not in the supplier list. Is this a new vendor or a wrong ID?", "Purchasing"),
        ("DQ-05", "RB-15M", "Two records exist for RB-15M ($1.31 and $1.29/kg). Which price is current?", "Purchasing"),
        ("DQ-07", "PC-MH-1200B", "BOM uses RB-12M, which does not exist. Should it be RB-10M or RB-15M?", "Engineering"),
        ("DQ-11", "PC-RCP-0900", "Rebar quantity is in LB, other BOMs use KG. Is 121 lb (about 55 kg) correct?", "Engineering"),
        ("DQ-01", "PC-HC-250", "HS code is missing for the 250 mm hollow-core slab. Confirm the classification.", "Product team"),
        ("DQ-14", "PC-RCP-0750", "Invoice lists item PC-RCP-0750, which is not in the item list. New product or typo?", "Sales"),
        ("DQ-15", "(2 invoices)", "Two invoices are priced 30-38% below list price. Were these discounts approved?", "Sales"),
        ("DQ-13", "(2 invoices)", "Two invoice lines have no unit price. Can the price be confirmed from the ERP?", "Accounting"),
    ]
    for i, (chk, ref, q, ask) in enumerate(questions):
        r = 5 + i
        put(ws, r, 1, f"Q{i+1:02d}", F_BOLD)
        put(ws, r, 2, raised, F_INPUT, "yyyy-mm-dd")
        put(ws, r, 3, chk, F_INPUT)
        put(ws, r, 4, ref, F_INPUT)
        c = put(ws, r, 5, q, F_INPUT); c.alignment = Alignment(wrap_text=True, vertical="top")
        put(ws, r, 6, ask, F_INPUT)
        put(ws, r, 7, "Open", F_INPUT)
        put(ws, r, 8, None, F_INPUT)
    oq_last = 4 + len(questions)
    ws.conditional_formatting.add(f"G5:G{oq_last+50}",
                                  CellIsRule(operator="equal", formula=['"Open"'], fill=FILL_FLAG))
    ws.conditional_formatting.add(f"G5:G{oq_last+50}",
                                  CellIsRule(operator="equal", formula=['"Closed"'], fill=FILL_OK))
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ About
    ws = ws_about
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    ws["A1"] = "Precast Cost and Sourcing Analysis"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Learning project with synthetic data for a fictional precast concrete manufacturer, "
                "Northfield Precast Ltd. No real company data is used.")
    ws["A2"].font = F_SUB
    rows = [
        ("Purpose", "Connect products, bills of materials, suppliers and sales to answer: what does each "
                    "product cost to make, where does the material come from, and which data needs fixing?"),
        ("", ""),
        ("Sheet", "What it contains"),
        ("Assumptions", "Exchange rates, unit conversions, thresholds and key assumptions (inputs)."),
        ("Items", "Finished products: list price, conversion cost, HS code. Raw data + checks."),
        ("Suppliers", "Supplier list with country and lead time."),
        ("Components", "Purchased materials and parts with cost, supplier and origin. Raw data + checks."),
        ("BOM", "Bill of materials lines. Looks up cost and origin, converts units, flags issues."),
        ("Sales", "2025 invoice lines. Adds category, revenue, discount and cost of sales."),
        ("Cost_Rollup", "Unit cost per product vs. average selling price: margin and cost-to-price ratio."),
        ("Sourcing", "2025 material spend by origin country and by supplier, with reconciliation checks."),
        ("Summary", "Revenue, cost of sales and margin by category and ship-to country."),
        ("Data_Quality", "Count of every flagged issue, with severity and a suggested action."),
        ("Open_Questions", "Log of questions to resolve with other teams."),
        ("", ""),
        ("Colour guide", ""),
        ("Blue text", "Hardcoded input (raw data or an assumption)."),
        ("Black text", "Formula."),
        ("Green text", "Direct link to another sheet."),
        ("Yellow fill", "Key assumption you can change on the Assumptions sheet."),
        ("Dark blue header", "Raw data column (as received)."),
        ("Grey-blue header", "Calculated column added in this workbook."),
        ("Red / green fill", "Status CHECK (needs review) / OK."),
        ("", ""),
        ("Rebuild", "python scripts/generate_data.py  then  python scripts/build_workbook.py"),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ca = ws.cell(row=i, column=1, value=a)
        cb = ws.cell(row=i, column=2, value=b)
        ca.font = F_BOLD if a else F_BASE
        cb.font = F_BASE
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if a in ("Sheet",):
            ca.font = cb.font = F_HEAD
            ca.fill = cb.fill = FILL_RAW
    legend_fonts = {"Blue text": F_INPUT, "Green text": F_LINK}
    for r in range(4, 4 + len(rows)):
        v = ws.cell(row=r, column=1).value
        if v in legend_fonts:
            ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True, color=legend_fonts[v].color)
        if v == "Yellow fill":
            ws.cell(row=r, column=1).fill = FILL_ASSUME
        if v == "Dark blue header":
            ws.cell(row=r, column=1).fill = FILL_RAW; ws.cell(row=r, column=1).font = F_HEAD
        if v == "Grey-blue header":
            ws.cell(row=r, column=1).fill = FILL_CALC; ws.cell(row=r, column=1).font = F_HEAD

    # Analysis sheets read like reports: no gridlines
    for w in (ws_about, ws_as, ws_cr, ws_so, ws_sm, ws_dq, ws_oq):
        w.sheet_view.showGridLines = False

    # Tab colours: inputs dark, analysis lighter
    for w in (ws_it, ws_su, ws_co, ws_bo, ws_sa):
        w.sheet_properties.tabColor = "1F3A5F"
    for w in (ws_cr, ws_so, ws_sm):
        w.sheet_properties.tabColor = "5B7F95"
    for w in (ws_dq, ws_oq):
        w.sheet_properties.tabColor = "C55A11"
    ws_as.sheet_properties.tabColor = "FFC000"

    # Print setup: landscape, fit all columns on one page width
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True

    if not compat:
        for w in wb.worksheets:
            for row in w.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        c.value = to_xlookup(c.value)

    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build workbook/Precast_Cost_Analysis.xlsx from data/raw/.")
    parser.add_argument("--compat", action="store_true", help="use INDEX/MATCH instead of XLOOKUP")
    build(compat=parser.parse_args().compat)
